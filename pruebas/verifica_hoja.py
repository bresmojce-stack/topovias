# Recalcula TODAS las fórmulas de cada hoja desde sus celdas y las compara con el valor que dejó la app
import openpyxl, re, sys, warnings
warnings.filterwarnings("ignore")
def evaluar(path):
    wb = openpyxl.load_workbook(path); wv = openpyxl.load_workbook(path, data_only=True)
    malos = total = 0
    for ws in wb.worksheets:
        wsv = wv[ws.title]; cache = {}
        def val(r, hoja=None):
            h = wb[hoja] if hoja else ws
            r = r.replace('$', '')
            key = (h.title, r)
            if key in cache: return cache[key]
            v = h[r].value
            if isinstance(v, str) and v.startswith('='):
                v = calc(v[1:], h.title)
            elif not isinstance(v, (int, float)): v = None if v in (None, '') else v
            cache[key] = v; return v
        def rango(a, b, hoja):
            ca, ra = re.match(r'\$?([A-Z]+)\$?(\d+)', a).groups(); cb, rb = re.match(r'\$?([A-Z]+)\$?(\d+)', b).groups()
            ci = lambda c: openpyxl.utils.column_index_from_string(c)
            out = []
            for rr in range(int(ra), int(rb) + 1):
                for cc in range(ci(ca), ci(cb) + 1): out.append(val(openpyxl.utils.get_column_letter(cc) + str(rr), hoja))
            return [x for x in out if isinstance(x, (int, float))]
        def rangoT(a, b, hoja):
            ca, ra = re.match(r'\$?([A-Z]+)\$?(\d+)', a).groups(); cb, rb = re.match(r'\$?([A-Z]+)\$?(\d+)', b).groups()
            ci = lambda c: openpyxl.utils.column_index_from_string(c)
            return [val(openpyxl.utils.get_column_letter(cc) + str(rr), hoja) for rr in range(int(ra), int(rb) + 1) for cc in range(ci(ca), ci(cb) + 1)]
        def calc(e, hoja):
            e = re.sub(r"'((?:[^']|'')+)'!(\$?[A-Z]+\$?\d+)", lambda m: f'V("{m.group(2)}","{m.group(1)}")', e)
            e = re.sub(r'SUMIF\((\$?[A-Z]+\$?\d+):(\$?[A-Z]+\$?\d+),(\$?[A-Z]+\$?\d+),(\$?[A-Z]+\$?\d+):(\$?[A-Z]+\$?\d+)\)',
                       lambda m: f'SUMIF_(RT("{m.group(1)}","{m.group(2)}"),VT("{m.group(3)}"),RT("{m.group(4)}","{m.group(5)}"))', e)
            e = re.sub(r'(SUM|AVERAGE)\((\$?[A-Z]+\$?\d+):(\$?[A-Z]+\$?\d+)\)', lambda m: f'{m.group(1)}(RG("{m.group(2)}","{m.group(3)}"))', e)
            e = re.sub(r'(?<![A-Za-z"!])(\$?[A-Z]{1,2}\$?\d+)\b(?!")', lambda m: f'V("{m.group(1)}")', e)
            env = {'V': lambda r, h=hoja: (val(r, h) or 0), 'VT': lambda r, h=hoja: ws[r.replace('$','')].value, 'RT': lambda a, b: rangoT(a, b, hoja),
                   'SUMIF_': lambda c, k, v: sum((y or 0) for x, y in zip(c, v) if x == k), 'RG': lambda a, b: rango(a, b, hoja), 'SUM': lambda *a: sum(sum(x) if isinstance(x, list) else x for x in a),
                   'AVERAGE': lambda *a: (lambda l: sum(l) / len(l) if l else 0)([y for x in a for y in (x if isinstance(x, list) else [x])]), 'ROUND': lambda x, n: float(__import__('decimal').Decimal(repr(round(x, 9))).quantize(__import__('decimal').Decimal(1).scaleb(-n), rounding=__import__('decimal').ROUND_HALF_UP)),
                   'MAX': max, 'ABS': abs, 'IF': lambda c, a, b: a if c else b}
            return eval(e.replace('^', '**'), env)
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith('='):
                    total += 1
                    try: ev = val(c.coordinate)
                    except Exception as ex: print('  ERROR', ws.title, c.coordinate, c.value, ex); malos += 1; continue
                    app = wsv[c.coordinate].value
                    if not isinstance(app, (int, float)) or ev is None or abs(ev - app) > 0.0006:
                        malos += 1
                        if malos < 15: print('  ✗', ws.title, c.coordinate, c.value, 'fórmula=', ev, 'app=', app)
    print(f'  {total} fórmulas, {malos} distintas')
    return malos
tot = 0
for p in sys.argv[1:]:
    print(p.split('/')[-1]); tot += evaluar(p)
print('PROBLEMAS:', tot)
