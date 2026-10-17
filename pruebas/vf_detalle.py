# Recalcula cada fórmula del Excel a partir de las celdas que referencia (sin usar los valores guardados)
import openpyxl, re, sys, warnings
warnings.filterwarnings("ignore")
def evaluar(path):
    wb = openpyxl.load_workbook(path); wv = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active; wsv = wv.active; cache = {}
    def val(ref):
        if ref in cache: return cache[ref]
        v = ws[ref].value
        if isinstance(v, str) and v.startswith('='):
            e = v[1:]
            e = re.sub(r'SUM\(([A-Z]+)(\d+):([A-Z]+)(\d+)\)', lambda m: '(' + '+'.join(f'val("{m.group(1)}{r}")' for r in range(int(m.group(2)), int(m.group(4)) + 1)) + ')', e)
            e = re.sub(r'(?<![A-Za-z"])([A-Z]{1,2}\d+)\b', lambda m: f'val("{m.group(1)}")', e)
            e = e.replace('IF(', '_if(').replace('ABS(', 'abs(')
            v = eval(e, {'val': val, '_if': lambda c, a, b: a if c else b, 'abs': abs})
        elif not isinstance(v, (int, float)): v = 0 if v in (None, '') else v
        cache[ref] = v; return v
    # sección de cantidades
    filas = [r for r in ws.iter_rows() if r[0].value == 'ITEM']
    assert filas, 'no hay sección ITEM'
    r0 = filas[0][0].row
    print('  encabezado:', [ws.cell(r0, k).value for k in range(1, 5)])
    malos = 0; r = r0 + 1
    while ws.cell(r, 1).value not in (None, ''):
        nom, uni, f = ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value
        app = wsv.cell(r, 3).value
        ev = val(f'C{r}')
        es_f = isinstance(f, str) and f.startswith('=')
        ok = abs(ev - app) <= 0.0006
        malos += (not ok) or (not es_f and nom != 'Área conformada')
        print(f'  {nom:16s} {uni:3s} app={app!r}  fórmula={ev!r}  {"✓" if ok else "✗"}  {"fórmula" if es_f else "VALOR FIJO"}: {str(f)[:110]}')
        r += 1
    return malos
tot = 0
for p in sys.argv[1:]:
    print(p.split('/')[-1]); tot += evaluar(p)
print('PROBLEMAS:', tot)
